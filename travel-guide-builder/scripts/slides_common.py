# -*- coding: utf-8 -*-
"""方案一（详细版）/ 方案二（简洁版）共用的骨架与渲染管线。

包含：
  · 路径常量、C 列骨架压缩
  · 两端共用的 6 个页面：封面 / 总行程 / 人均预算 / 出发前必看 / 物品清单 / 结束祝福
  · 「Jinja2 → Chrome headless 截图 → python-pptx」的渲染导出管线（含自动缩排兜底）

两个方案只有**中间的逐日页面**不同：
  方案一 每景点 2 页（详细 + 美景机位）+ 当天行程页 + 当天收尾页
  方案二 每天 2 页（行程规划 + 餐饮住宿）
"""
import os
import re
import json

from jinja2 import Environment, FileSystemLoader
from PIL import Image
from pptx import Presentation
from pptx.util import Inches
from playwright.sync_api import sync_playwright

BASE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.join(BASE, "roadbook")
RAW = os.path.join(ROOT, "_raw", "detail")
SLIDES_JSON = os.path.join(RAW, "slides.json")
PHOTO_MAP = os.path.join(RAW, "photo_map.json")
PPT = os.path.join(ROOT, "assets", "ppt")
MAPS = os.path.join(ROOT, "assets", "maps_xhs")
CHROME = r"C:\Program Files\Google\Chrome\Application\chrome.exe"

COVER_IMG = "huangguoshu/01.jpg"
ROUTE_NAME = "贵阳环线 · 黔中铁三角"
MILES = "约 950 km"

# 交通过渡行（压缩总行程骨架时丢弃）
TRANSIT_RE = re.compile(r"^[✈🚗🚶🚌🍜🛒🧳⛽🚕]\s*")
TIME_RE = re.compile(r"\d+\s*(min|h|小时|分钟)")

# 每天的城市景点地图
DAY_MAP = {
    1: ("guiyang_city", "贵阳市景点分布"),
    2: ("anshun_city", "安顺景点分布"),
    3: ("libo_city", "荔波旅游地图"),
    4: ("libo_city", "荔波旅游地图"),
    5: ("qdn_city", "黔东南旅游地图"),
    6: ("guiyang_city", "贵阳市景点分布"),
    7: ("guiyang_city", "贵阳市景点分布"),
}


def U(p):
    """本地路径 → file:/// URL（给 Chrome 用）"""
    return "file:///" + os.path.abspath(p).replace("\\", "/")


def read_json(p):
    with open(p, encoding="utf-8") as f:
        return json.load(f)


# ------------------------------------------------------------------ 骨架压缩

def condense_stop(line):
    """C 列一行 → 骨架用的站点名；返回 None 表示这是交通过渡行，丢弃。"""
    s = line.strip()
    if not s:
        return None
    is_transit = bool(TRANSIT_RE.match(s))
    if is_transit:
        body = TRANSIT_RE.sub("", s)
        # 带明确时长的交通行一律丢
        if TIME_RE.search(body) or len(body) <= 8:
            return None
    s = re.sub(r"（[^）]*）", "", s)
    s = re.sub(r"\([^)]*\)", "", s)
    s = TRANSIT_RE.sub("", s)
    s = s.rstrip("→ ").strip()
    return s or None


def build_chain(brief):
    stops = []
    for ln in brief:
        c = condense_stop(ln)
        if c and (not stops or stops[-1] != c):
            stops.append(c)
    return stops


# ------------------------------------------------------------------ 共用页面

def cover_slide(meta):
    """01 封面：某目的地 · 某天某晚 · 旅游攻略"""
    return {
        "kind": "cover",
        "cover_img": U(os.path.join(PPT, COVER_IMG)),
        "title_pre": "%s " % meta["dest"],
        "title_hl": "%d 天 %d 晚" % (meta["days"], meta["nights"]),
        "title_post": "自驾旅游攻略",
        "dateRange": meta["dateRange"],
        "route": ROUTE_NAME,
        "miles": MILES,
        "chips": ["✈ 深圳 ⇄ 贵阳", "🚗 落地自驾", "👥 2 人",
                  "🏨 经济型", "⏰ 舒适型 · 每日 9:00 出发"],
    }


def overview_slide(days):
    """02 总行程：全部行的 C 列骨架 + 总路线地图"""
    ov = []
    for dy in days:
        ov.append({
            "label": "DAY%d · %s %s" % (dy["day"], dy["date"], dy["weekday"]),
            "main": dy["transport_main"],
            "chain": " <span class='arw'>→</span> ".join(build_chain(dy["brief"])),
        })
    return {
        "kind": "overview",
        "map": U(os.path.join(MAPS, "zongluxian.png")),
        "map_cap": "贵州黔中环线 · 总路线图（约 950 km，日均车程 2-3h）",
        "days": ov,
    }


def budget_slide(data):
    """03 人均预算"""
    b = data["budget"]
    tot = b.get("totals", [])
    wk = ho = wk2 = ho2 = ""
    for row in tot:
        if any("平日" in x and "¥" in x for x in row):
            wk, ho = row[2], row[3]
            wk2 = "2 人合计 " + (row[5].split("：")[1].split("/")[0].strip()
                              if "：" in row[5] else "")
            ho2 = "2 人合计 " + (row[5].split("/")[-1].strip()
                              if "/" in row[5] else "")
    warn = ""
    for row in tot:
        if row[0].startswith("⚠️"):
            warn = row[0]
    return {
        "kind": "budget",
        "headers": b["headers"], "rows": b["rows"],
        "weekday": wk, "holiday": ho,
        "weekday2": wk2, "holiday2": ho2,
        "warning": warn,
    }


def preflight_slide(data):
    """04 出发前必看"""
    pf = data["preflight"]
    return {"kind": "preflight",
            "book": pf["book"], "avoid": pf["avoid"], "wear": pf["wear"]}


def packing_slide(data):
    """次尾页 物品清单"""
    # 注意：不能叫 items —— Jinja2 里 dict.items 是内置方法，会撞名
    return {"kind": "packing",
            "groups": [{"head": g["head"], "rows": g["items"]}
                       for g in data["packing"]]}


def closing_slide(meta, days):
    """尾页 结束祝福"""
    n_spot = sum(len(dy["spots"]) for dy in days)
    return {
        "kind": "closing",
        "cover_img": U(os.path.join(PPT, COVER_IMG)),
        "title_pre": "", "title_hl": "旅途顺利", "title_post": "",
        "wish": "祝大家旅行顺利 · 满载而归",
        "sub": "%s · %s · %s" % (meta["dateRange"], ROUTE_NAME, MILES),
        "chips": ["🚗 自驾 %d 天 %d 晚" % (meta["days"], meta["nights"]),
                  "📸 %d 个景点" % n_spot,
                  "🍲 贵州味道", "🛍 手信满载"],
    }


def head_slides(meta, days, data):
    """两个方案共用的前 4 页。"""
    return [cover_slide(meta), overview_slide(days),
            budget_slide(data), preflight_slide(data)]


def tail_slides(meta, days, data):
    """两个方案共用的末 2 页：次尾页物品清单 + 尾页结束祝福。"""
    return [packing_slide(data), closing_slide(meta, days)]


# ------------------------------------------------------------------ 去重
# 同一句提示经常同时出现在当日总备注（H 列）和某个非景点块（🚉 交通 / 🛍 事务）里 ——
# 例：DAY1「10 月初贵阳 18-26℃…」既是当日穿搭、又是龙洞堡机场节点的穿搭。
# 不做**跨块**去重就会在同一页出现两遍，用户侧看起来就是「内容有点乱」。

def _sk(s):
    return re.sub(r"\s+", "", s or "")


def dedup_lines(seq):
    """按内容去重（保序，忽略空白差异）。"""
    seen, out = set(), []
    for x in seq or []:
        k = _sk(x)
        if k and k not in seen:
            seen.add(k)
            out.append(x)
    return out


def dedup_blocks(blocks):
    """跨块去重：后面的块若与前面重复则丢弃该行；整块被抽空则丢掉这个块。"""
    seen, out = set(), []
    for b in blocks or []:
        lines = []
        for x in b.get("lines", []):
            k = _sk(x)
            if k and k not in seen:
                seen.add(k)
                lines.append(x)
        if lines:
            out.append({"head": b.get("head", ""), "lines": lines})
    return out


# ------------------------------------------------------------------ 渲染导出

AUTOFIT_JS = """() => {
    const out = [];
    document.querySelectorAll('.slide').forEach((s, i) => {
        const sb = s.getBoundingClientRect().bottom;
        s.querySelectorAll('.card-b, .fit').forEach(b => {
            const base = parseFloat(getComputedStyle(b).fontSize);
            const over = () => Math.max(
                b.scrollHeight - b.clientHeight,
                Math.ceil(b.getBoundingClientRect().bottom - (sb - 6)));
            if (over() <= 0) return;
            let size = base, guard = 0;
            while (over() > 0 && size > 12 && guard++ < 80) {
                size = Math.max(12, size - 0.5);
                const k = size / base;
                b.style.fontSize = size + 'px';
                b.style.lineHeight = '1.44';
                const kids = b.querySelectorAll('p, li');
                kids.forEach((e, n) => {
                    e.style.marginBottom = (n === kids.length - 1) ? '0'
                        : Math.max(1, Math.round(8 * k)) + 'px';
                });
                b.querySelectorAll('.mini').forEach(e => {
                    e.style.fontSize = (16.5 * k).toFixed(1) + 'px';
                    e.style.margin = Math.max(3, Math.round(10 * k)) + 'px 0 '
                                   + Math.max(1, Math.round(4 * k)) + 'px';
                });
            }
            if (size < base) out.push({page: i + 1, from: base, to: size, over: over()});
        });
    });
    return out;
}"""


def render_and_export(slides, tpl_path, work, out_html, out_pptx,
                      prefix="d", fitted_name=None, clean=True):
    """Jinja2 → Chrome 逐页截图 → pptx。返回 (shots, out_fitted)。

    prefix      : 截图文件名前缀（方案一 d / 方案二 c），避免两套产物互相覆盖
    clean       : 构建前清掉同前缀旧帧（否则页数变少时会残留旧页）
    """
    import glob

    print("slides =", len(slides))
    from collections import Counter
    print("页型分布:", dict(Counter(s["kind"] for s in slides)))

    # 用 Environment + FileSystemLoader（而不是 Template(str)）——
    # 只有带 loader 的 Environment 才支持 {% include %}，
    # 两套模板共用的 `_ppt_base.css.j2` 正是靠 include 引入的。
    tpl_dir, tpl_name = os.path.split(os.path.abspath(tpl_path))
    env = Environment(loader=FileSystemLoader(tpl_dir), trim_blocks=False,
                      lstrip_blocks=False)
    html = env.get_template(tpl_name).render(slides=slides)
    with open(out_html, "w", encoding="utf-8") as f:
        f.write(html)
    print("HTML:", out_html, len(html), "bytes")

    if clean:
        # ⚠️ 只删「序号超出本次页数」的旧帧 —— 页数变少时才需要清（否则残留页会混进
        #    缩略图/验收图）。不要无脑 rm d_*.png：一来本轮帧可以直接覆盖，二来
        #    一次性删几十个文件会被宿主的批量删除保护拦下并直接终止构建进程。
        stale = []
        for p in (glob.glob(os.path.join(work, prefix + "_*.png")) +
                  glob.glob(os.path.join(work, prefix + "_*.jpg"))):
            m = re.search(r"_(\d+)\.[A-Za-z]+$", os.path.basename(p))
            if m and int(m.group(1)) > len(slides):
                stale.append(p)
        for p in stale:
            os.remove(p)
        print("清理超范围旧帧:", len(stale))

    shots = []
    out_fitted = os.path.join(work, fitted_name) if fitted_name else None
    with sync_playwright() as p:
        br = p.chromium.launch(executable_path=CHROME, headless=True,
                               args=["--no-sandbox", "--font-render-hinting=none"])
        pg = br.new_page(viewport={"width": 1920, "height": 1080}, device_scale_factor=2)
        pg.goto("file:///" + out_html.replace("\\", "/"), wait_until="load")
        # 等所有图片真正解码完成，避免截到空白框
        pg.wait_for_timeout(1500)
        try:
            pg.wait_for_function(
                "() => Array.from(document.images).every(i => i.complete && i.naturalWidth > 0)",
                timeout=60000)
        except Exception as e:
            bad = pg.evaluate("""() => Array.from(document.images)
                .filter(i => !(i.complete && i.naturalWidth > 0))
                .map(i => i.getAttribute('src')).slice(0,10)""")
            print("!! 有图未加载:", bad, str(e)[:120])
        pg.wait_for_timeout(900)

        # ── 自动缩排兜底：内容装不下时按 0.5px 步长收字号 ──
        fitted = pg.evaluate(AUTOFIT_JS)
        if fitted:
            print("自动缩排（原内容装不下，已收字号）:")
            for f in fitted:
                flag = "" if f["over"] <= 0 else "  ⚠️ 仍差 %dpx，需人工删减内容" % f["over"]
                print("   P%02d  %.1fpx → %.1fpx%s"
                      % (f["page"], f["from"], f["to"], flag))
        else:
            print("自动缩排: 无需收缩，页面全部装得下")
        pg.wait_for_timeout(400)

        if out_fitted:
            snap = pg.evaluate(
                "() => '<!DOCTYPE html>\\n' + document.documentElement.outerHTML")
            with open(out_fitted, "w", encoding="utf-8") as f:
                f.write(snap)
            print("已导出缩排后快照:", os.path.basename(out_fitted))

        n = pg.locator(".slide").count()
        print("slides in html =", n)
        for i in range(n):
            fp = os.path.join(work, "%s_%02d.png" % (prefix, i + 1))
            pg.locator(".slide").nth(i).screenshot(path=fp)
            shots.append(fp)
        br.close()

    prs = Presentation()
    prs.slide_width = Inches(13.3333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]
    for fp in shots:
        s = prs.slides.add_slide(blank)
        im = Image.open(fp).convert("RGB")
        if im.size != (3840, 2160):
            im = im.resize((3840, 2160), Image.LANCZOS)
        jp = fp.replace(".png", ".jpg")
        im.save(jp, "JPEG", quality=90, optimize=True)
        s.shapes.add_picture(jp, 0, 0, width=prs.slide_width, height=prs.slide_height)
    prs.save(out_pptx)
    print("PPTX:", out_pptx, round(os.path.getsize(out_pptx) / 1024), "KB")
    return shots, out_fitted
