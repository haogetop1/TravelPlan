# -*- coding: utf-8 -*-
"""从同一份 content.json 生成 PPT 路书（与 HTML 内容一致）。

用法：
  python build_ppt.py             # 阶段1：封面 + 总览 + DAY1
  python build_ppt.py --stage 7   # 全量
"""
import os, re, sys, json, html, argparse, hashlib
from PIL import Image
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

BASE = os.path.dirname(os.path.abspath(__file__))
RB = os.path.join(BASE, 'roadbook')
PHOTO_SRC = os.path.join(RB, '_raw', 'photos')
MAPS = os.path.join(RB, 'assets', 'maps')
CROP = os.path.join(RB, '_raw', 'pptcrop')
os.makedirs(CROP, exist_ok=True)

sys.path.insert(0, BASE)
from build_html import PICKS, HERO, ESSENTIALS, load_meta, all_photos, clean_line, norm_name  # noqa

GOLD = RGBColor(0xB0, 0x89, 0x68)
GOLD_D = RGBColor(0x8A, 0x6A, 0x4F)
GOLD_BG = RGBColor(0xF3, 0xEC, 0xE4)
CREAM = RGBColor(0xFA, 0xF7, 0xF2)
INK = RGBColor(0x2F, 0x2A, 0x26)
INK2 = RGBColor(0x4A, 0x42, 0x3B)
MUTED = RGBColor(0x85, 0x7A, 0x70)
LINE = RGBColor(0xE8, 0xE0, 0xD5)
RED = RGBColor(0xB2, 0x3B, 0x3B)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
DARKBG = RGBColor(0x24, 0x1D, 0x17)

FONT = '微软雅黑'
SW, SH = 13.333, 7.5


def blank(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])


def bg(slide, color=CREAM):
    sh = slide.shapes.add_shape(1, 0, 0, Inches(SW), Inches(SH))
    sh.fill.solid(); sh.fill.fore_color.rgb = color
    sh.line.fill.background(); sh.shadow.inherit = False
    return sh


def rect(slide, x, y, w, h, fill, line=None, lw=0.75):
    sh = slide.shapes.add_shape(1, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill is None:
        sh.fill.background()
    else:
        sh.fill.solid(); sh.fill.fore_color.rgb = fill
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line; sh.line.width = Pt(lw)
    sh.shadow.inherit = False
    return sh


def tx(slide, x, y, w, h, runs, size=14, bold=False, color=INK,
       align=PP_ALIGN.LEFT, spacing=1.25, anchor=MSO_ANCHOR.TOP, wrap=True):
    """runs: str 或 [(text, {size,bold,color}), ...]（每项一段）"""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = wrap
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    if isinstance(runs, str):
        runs = [(runs, {})]
    for i, (t, opt) in enumerate(runs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = opt.get('align', align)
        p.line_spacing = opt.get('spacing', spacing)
        if opt.get('space_before'):
            p.space_before = Pt(opt['space_before'])
        r = p.add_run(); r.text = t
        f = r.font
        f.name = FONT
        f.size = Pt(opt.get('size', size))
        f.bold = opt.get('bold', bold)
        f.color.rgb = opt.get('color', color)
    return tb


def pic_cover(slide, path, x, y, w, h):
    """按 cover 方式裁剪填充目标框（居中裁剪）。"""
    try:
        im = Image.open(path)
    except Exception:
        return None
    tw, th = w, h
    ar_t = tw / th
    ar_s = im.size[0] / im.size[1]
    key = hashlib.md5(('%s|%.2f' % (path, ar_t)).encode('utf-8')).hexdigest()[:12]
    out = os.path.join(CROP, '%s.jpg' % key)
    if not os.path.exists(out):
        if ar_s > ar_t:      # 源更宽 → 裁左右
            nw = int(im.size[1] * ar_t)
            left = (im.size[0] - nw) // 2
            im = im.crop((left, 0, left + nw, im.size[1]))
        else:                # 源更高 → 裁上下
            nh = int(im.size[0] / ar_t)
            top = int((im.size[1] - nh) * 0.42)
            im = im.crop((0, top, im.size[0], top + nh))
        if im.size[0] > 1600:
            im = im.resize((1600, int(im.size[1] * 1600 / im.size[0])), Image.LANCZOS)
        im.convert('RGB').save(out, 'JPEG', quality=86, optimize=True)
    return slide.shapes.add_picture(out, Inches(x), Inches(y), Inches(w), Inches(h))


def pic_fit(slide, path, x, y, w, h):
    """按 contain 方式放进框内并居中。"""
    try:
        im = Image.open(path)
    except Exception:
        return None
    ar_t, ar_s = w / h, im.size[0] / im.size[1]
    if ar_s > ar_t:
        nw, nh = w, w / ar_s
    else:
        nh, nw = h, h * ar_s
    return slide.shapes.add_picture(path, Inches(x + (w - nw) / 2), Inches(y + (h - nh) / 2),
                                   Inches(nw), Inches(nh))


def header(slide, kicker, title, en='', color=INK):
    rect(slide, 0.55, 0.42, 0.055, 0.52, GOLD)
    tx(slide, 0.78, 0.40, 9.0, 0.30, kicker, size=11.5, bold=True, color=GOLD_D)
    tx(slide, 0.78, 0.62, 9.6, 0.42, title, size=23, bold=True, color=color)
    if en:
        tx(slide, 8.6, 0.45, 4.18, 0.3, en, size=11, color=MUTED, align=PP_ALIGN.RIGHT)
    rect(slide, 0.55, 1.13, SW - 1.1, 0.022, GOLD)


def bullets(runs_src, limit=4, width=58):
    out = []
    for s in runs_src:
        s = re.sub(r'\s+', ' ', (s or '').strip())
        s = re.sub(r'^[·\-•]\s*', '', s)
        if not s:
            continue
        out.append(s[:width] + ('…' if len(s) > width else ''))
        if len(out) >= limit:
            break
    return out


def wrap_cn(s, width):
    s = re.sub(r'\s+', ' ', (s or '').strip())
    s = re.sub(r'^[·\-•]\s*', '', s)
    return s[:width] + ('…' if len(s) > width else '')


def chain_text(brief):
    return '  →  '.join(clean_line(x) for x in brief if clean_line(x))


def sec_kind(name):
    if any(k in name for k in ('机场', '登机', '还车')):
        return 'airport'
    if any(k in name for k in ('市区', '县城', '采购', '步行街', '小镇')):
        return 'city'
    return 'spot'


def build(stage):
    content = json.load(open(os.path.join(RB, 'content.json'), encoding='utf-8'))
    meta = content['meta']
    b = content['budget']
    ph_meta = load_meta()
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(SW), Inches(SH)

    # ---------- 1. 封面 ----------
    s = blank(prs); bg(s, DARKBG)
    if HERO:
        items = all_photos(HERO[0])
        if 1 <= HERO[1] <= len(items):
            _, _, hp = items[HERO[1] - 1]
            pic_cover(s, hp, 0, 0, SW, SH)
    ov = rect(s, 0, 0, SW, SH, DARKBG)
    ov.fill.transparency = 0.34
    rect(s, 0, 4.05, SW, 3.45, DARKBG)
    rect(s, 0.9, 1.55, 0.9, 0.05, GOLD)
    tx(s, 0.9, 1.78, 8.0, 0.35, 'SELF-DRIVE ROADBOOK  ·  2026 AUTUMN', size=13, bold=True, color=RGBColor(0xE0, 0xC2, 0x9E))
    tx(s, 0.9, 2.22, 11.0, 1.95,
       [('贵州 %d 天 %d 晚' % (meta['days'], meta['nights']), {'size': 46, 'bold': True, 'color': WHITE, 'spacing': 1.05}),
        ('%s旅游攻略' % meta['mode'], {'size': 46, 'bold': True, 'color': RGBColor(0xE8, 0xC9, 0xA0), 'spacing': 1.05})])
    rect(s, 0.9, 4.42, 1.4, 0.04, GOLD)
    tx(s, 0.9, 4.66, 11.5, 0.4,
       '%s  ·  %s  ·  全程约 %s' % (meta['dateRange'], meta['route'], meta['miles']),
       size=15, color=RGBColor(0xD8, 0xCF, 0xC4))
    chips = ['✈ 深圳 ⇄ 贵阳', '🚗 落地自驾', '👥 %d 人' % meta['people'],
             '🏨 %s' % meta['level'], '⏰ %s' % meta['pace']]
    x = 0.9
    for c in chips:
        w = 0.34 + len(c) * 0.115
        rect(s, x, 5.32, w, 0.46, RGBColor(0x3A, 0x2F, 0x26), RGBColor(0x6A, 0x5A, 0x4A))
        tx(s, x, 5.42, w, 0.3, c, size=12, color=RGBColor(0xE8, 0xDF, 0xD5), align=PP_ALIGN.CENTER)
        x += w + 0.16
    tx(s, 0.9, 6.5, 11.5, 0.3, '行程 / 预算 / 地图 / 机位 / 餐饮 / 穿衣避坑 / 注意事项',
       size=12.5, color=RGBColor(0x9A, 0x8C, 0x7C))

    # ---------- 2. 路线总览 ----------
    s = blank(prs); bg(s)
    header(s, '总篇章', '7 天路线骨架', 'OVERVIEW')
    y = 1.42
    for d in content['days']:
        h = 0.72
        rect(s, 0.55, y, 7.55, h, WHITE, LINE)
        rect(s, 0.55, y, 0.055, h, GOLD)
        tx(s, 0.78, y + 0.09, 1.9, 0.3, 'DAY%d · %s %s' % (d['day'], d['date'], d['weekday']),
           size=12, bold=True, color=GOLD_D)
        tx(s, 2.62, y + 0.07, 3.3, 0.32, '%s → %s' % (d['from_city'], d.get('to_city') or d.get('city') or ''),
           size=13, bold=True, color=INK)
        tx(s, 0.78, y + 0.38, 7.1, 0.3, chain_text(d['brief'])[:78], size=9.5, color=MUTED)
        y += h + 0.075
    pic_fit(s, os.path.join(MAPS, 'zongluxian.png'), 8.35, 1.42, 4.42, 3.6)
    rect(s, 8.35, 5.12, 4.42, 2.0, WHITE, LINE)
    tx(s, 8.55, 5.28, 4.0, 0.3, '总路线地图', size=13.5, bold=True, color=INK)
    tx(s, 8.55, 5.62, 4.05, 1.4,
       [('贵阳 → 黄果树 → 荔波 → 西江 → 贵阳', {'size': 11.5, 'color': INK2}),
        ('全程约 %s，日均车程 2-3h' % meta['miles'], {'size': 11.5, 'color': INK2}),
        ('高德地图 · 审图号 GS(2025)5996 号', {'size': 10, 'color': MUTED, 'space_before': 5})])

    # ---------- 3. 人均预算 ----------
    s = blank(prs); bg(s)
    header(s, '总篇章', '人均总预算', 'BUDGET')
    rect(s, 0.55, 1.4, 4.1, 1.5, GOLD)
    tx(s, 0.82, 1.54, 3.6, 0.3, '平日参考价 · 人均', size=12, color=RGBColor(0xF3, 0xE8, 0xDC))
    tx(s, 0.82, 1.85, 3.6, 0.7, b['per_weekday'], size=34, bold=True, color=WHITE)
    tx(s, 0.82, 2.53, 3.6, 0.3, '2 人合计 %s' % b['total_weekday'], size=12, color=RGBColor(0xF3, 0xE8, 0xDC))
    rect(s, 4.85, 1.4, 4.1, 1.5, WHITE, LINE)
    tx(s, 5.12, 1.54, 3.6, 0.3, '国庆实际价 · 人均', size=12, color=MUTED)
    tx(s, 5.12, 1.85, 3.6, 0.7, b['per_holiday'], size=34, bold=True, color=RED)
    tx(s, 5.12, 2.53, 3.6, 0.3, '2 人合计 %s' % b['total_holiday'], size=12, color=MUTED)
    rect(s, 9.15, 1.4, 3.62, 1.5, GOLD_BG, LINE)
    tx(s, 9.4, 1.54, 3.2, 0.3, '国庆溢价', size=12, color=GOLD_D)
    tx(s, 9.4, 1.86, 3.2, 0.6, b.get('delta', '1.6×'), size=30, bold=True, color=GOLD_D)
    tx(s, 9.4, 2.5, 3.2, 0.3, '机票 2.8× / 住宿 2.0×', size=10.5, color=GOLD_D)

    y = 3.08
    rect(s, 0.55, y, SW - 1.1, 0.34, GOLD_BG)
    for cx, cw, t in [(0.75, 3.9, '类别'), (4.7, 3.4, '平日 / 人'), (8.15, 3.4, '国庆 / 人'), (11.6, 1.2, '涨幅')]:
        tx(s, cx, y + 0.055, cw, 0.26, t, size=11.5, bold=True, color=GOLD_D)
    y += 0.34
    for i, r in enumerate(b['rows']):
        h = 0.42
        if i % 2 == 0:
            rect(s, 0.55, y, SW - 1.1, h, RGBColor(0xFC, 0xFA, 0xF7))
        tx(s, 0.75, y + 0.08, 3.9, 0.28, r[0], size=12, bold=True, color=INK)
        tx(s, 4.7, y + 0.08, 3.4, 0.28, r[2], size=12, color=INK2)
        tx(s, 8.15, y + 0.08, 3.4, 0.28, r[3], size=12, bold=True, color=RED)
        tx(s, 11.6, y + 0.08, 1.2, 0.28, r[4], size=11.5, color=MUTED)
        y += h
    rect(s, 0.55, y + 0.06, SW - 1.1, 0.7, RGBColor(0xFD, 0xF6, 0xEC), GOLD)
    tx(s, 0.78, y + 0.17, SW - 1.6, 0.5,
       [('平日与国庆差距主要来自机票 2.8× 与住宿 2.0×；国庆高速免费，过路费 ¥0，油路费反而更省 ¥175/人。建议尽早锁定机票与西江住宿。', {'size': 10.5, 'color': RGBColor(0x6B, 0x5A, 0x48)})])

    # ---------- 4-... 每日 ----------
    for d in content['days']:
        if d['day'] > stage:
            continue
        # 日概览
        s = blank(prs); bg(s)
        header(s, 'DAY %d' % d['day'], '%s %s' % (d['date'], d['weekday']),
               '%s → %s' % (d['from_city'], d.get('to_city') or d.get('city') or ''))
        rect(s, 0.55, 1.4, 5.6, 1.55, WHITE, LINE)
        tx(s, 0.78, 1.54, 5.2, 0.28, '大交通', size=12, bold=True, color=GOLD_D)
        tx(s, 0.78, 1.85, 5.2, 0.95, (d.get('transport_raw') or '')[:80], size=12, color=INK2, spacing=1.4)
        rect(s, 6.35, 1.4, 6.42, 1.55, WHITE, LINE)
        tx(s, 6.58, 1.54, 6.0, 0.28, '住宿城市', size=12, bold=True, color=GOLD_D)
        tx(s, 6.58, 1.85, 6.0, 0.9, (d.get('stay_note') or d.get('city') or '') + '（最终住宿后续再定）',
           size=12, color=INK2, spacing=1.4)
        rect(s, 0.55, 3.08, 12.22, 1.5, WHITE, LINE)
        tx(s, 0.78, 3.2, 11.8, 0.28, '简要行程', size=12, bold=True, color=GOLD_D)
        tx(s, 0.78, 3.5, 11.8, 1.0, chain_text(d['brief']), size=11.5, color=INK2, spacing=1.55)
        if d.get('city_map'):
            pic_fit(s, os.path.join(MAPS, d['city_map'] + '.png'), 0.55, 4.72, 8.4, 2.45)
            rect(s, 0.55, 4.72, 8.4, 2.45, None, LINE)
            tx(s, 0.72, 6.92, 8.1, 0.26, '%s 景点分布图 · 高德地图 | 审图号 GS(2025)5996 号' % (d.get('city') or ''),
               size=9.5, color=MUTED)
        foods = [clean_line(x) for x in (d.get('food') or [])][:5]
        rect(s, 9.15, 4.72, 3.62, 2.45, WHITE, LINE)
        tx(s, 9.35, 4.84, 3.3, 0.28, '当日餐饮', size=12, bold=True, color=GOLD_D)
        tx(s, 9.35, 5.14, 3.32, 1.9,
           [(wrap_cn(x, 26), {'size': 9.5, 'color': INK2, 'space_before': 3}) for x in foods] or
           [('待补充', {'size': 9.5, 'color': MUTED})])

        # 每个大景点
        for sec in d['sections']:
            kind = sec_kind(sec['name'])
            if kind == 'airport' and not sec.get('map'):
                continue
            s = blank(prs); bg(s)
            tag = '交通枢纽' if kind == 'airport' else ('城市' if kind == 'city' else '大景点')
            header(s, '%s · %s' % (d['date'], tag), sec['short'], 'DAY %d' % d['day'], color=RED)

            # 上图 / 左图
            if sec.get('map'):
                pic_fit(s, os.path.join(MAPS, sec['map'] + '.png'), 0.55, 1.4, 5.6, 3.0)
                rect(s, 0.55, 1.4, 5.6, 3.0, None, LINE)
            # 详细行程
            detail = re.sub(r'\s*\n\s*', '　', sec.get('detail') or '')
            tx(s, 6.35, 1.42, 6.42, 3.0,
               [('详细行程', {'size': 12, 'bold': True, 'color': GOLD_D}),
                (detail[:430] + ('…' if len(detail) > 430 else ''), {'size': 9.5, 'color': INK2, 'space_before': 4})],
               spacing=1.32)

            # 照片
            picks = PICKS.get(sec.get('slug') or '', [])
            items = all_photos(sec.get('slug') or '')
            photos = []
            for idx, cap in picks:
                if 1 <= idx <= len(items):
                    photos.append((items[idx - 1][2], cap))
            photos = photos[:3]
            px, pw, ph_ = 0.55, 2.95, 2.2
            for i, (pp, cap) in enumerate(photos):
                pic_cover(s, pp, px + i * (pw + 0.16), 4.62, pw, ph_)
                tx(s, px + i * (pw + 0.16), 6.86, pw, 0.24, cap, size=8.5, color=MUTED)
            if not photos:
                rect(s, 0.55, 4.62, 9.3, ph_, GOLD_BG, LINE)
                tx(s, 0.75, 5.5, 9.0, 0.3, '（该点位暂无实拍配图）', size=11, color=MUTED, align=PP_ALIGN.CENTER)

            # 右下信息盒：避坑 / 注意事项 / 机位
            tips = bullets(sec.get('tips') or [], 4, 26)
            notice = bullets(sec.get('notice') or [], 4, 26)
            shots = []
            if sec.get('shots'):
                for blk in sec['shots']:
                    shots += bullets(blk.get('lines') or [], 3, 24)
            shots = shots[:3]
            boxx = 9.95
            rect(s, boxx, 4.62, 2.82, 2.2, WHITE, LINE)
            tx(s, boxx + 0.16, 4.72, 2.5, 0.26, '📸 机位', size=11, bold=True, color=GOLD_D)
            tx(s, boxx + 0.16, 5.0, 2.54, 1.7,
               [(x, {'size': 8.5, 'color': INK2, 'space_before': 3}) for x in (shots or ['—'])])

            if tips or notice:
                s2 = blank(prs); bg(s2)
                header(s2, '%s · %s' % (d['date'], sec['short']), '穿衣避坑 · 注意事项', 'DAY %d' % d['day'])
                rect(s2, 0.55, 1.4, 6.0, 5.6, WHITE, LINE)
                tx(s2, 0.8, 1.55, 5.5, 0.3, '🧥 备注 · 穿衣与避坑', size=13, bold=True, color=GOLD_D)
                tx(s2, 0.8, 1.92, 5.55, 5.0,
                   [(wrap_cn(x, 30), {'size': 10, 'color': INK2, 'space_before': 6}) for x in bullets(sec.get('tips') or [], 10, 30)])
                rect(s2, 6.82, 1.4, 5.95, 5.6, WHITE, LINE)
                tx(s2, 7.05, 1.55, 5.5, 0.3, '⚠️ 注意事项', size=13, bold=True, color=GOLD_D)
                tx(s2, 7.05, 1.92, 5.5, 5.0,
                   [(wrap_cn(x, 30), {'size': 10, 'color': INK2, 'space_before': 6}) for x in bullets(sec.get('notice') or [], 10, 30)])

    out = os.path.join(RB, '贵州7天6晚自驾路书.pptx')
    prs.save(out)
    print('written', out, len(prs.slides.__iter__.__self__._sldIdLst), 'slides')
    return out


if __name__ == '__main__':
    ap = argparse.ArgumentParser()
    ap.add_argument('--stage', type=int, default=1)
    a = ap.parse_args()
    build(a.stage)
