# -*- coding: utf-8 -*-
"""从攻略 xlsx 抽出结构化 content.json（路书唯一数据源）。

按【大景点】切块：D/E/H/I 四列都用同一套景点名做键，逐日对齐。
"""
import os, re, json, sys
from openpyxl import load_workbook

sys.stdout.reconfigure(encoding='utf-8')

# ── 非景点块判定统一走共享模块（HTML 链路与 PPT 链路必须引用同一份）──
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from nonspot_rules import split_blocks, clean_items, KIND_ICON  # noqa: E402

TAIL_RE = re.compile(r'^\s*📌\s*(当日总备注|当日总注意事项|总备注|总注意事项)\s*[：:]?\s*')


def split_tail(body):
    """把误并入景点块尾部的「📌 当日总备注 / 总注意事项」切出来。

    它作为无标记的行落在最后一个块里，不切出来会跟景点正文一起显示。
    """
    tail, keep = [], []
    for l in lines_of(body):
        (tail if TAIL_RE.match(l) else keep).append(l)
    return '\n'.join(keep), tail


def tail_block(t):
    """「📌 当日总备注：xxx」→ dict(name=小标题, lines=[正文])"""
    parts = re.split(r'[：:]', re.sub(r'^\s*📌\s*', '', str(t)), 1)
    return dict(name='📌 ' + parts[0].strip(),
                lines=[parts[1].strip()] if len(parts) > 1 and parts[1].strip() else [])


def block_label(name, kind):
    """非景点块在当日备注 / 注意事项里的小标题（🚉 交通节点 / 🛍 事务节点）。"""
    return '%s %s' % (KIND_ICON.get(kind, '📌'), (name or '当日说明').strip())


# 名字里带这些词的 `【】` 块其实不是景点（如「【敦煌抵达提示】」「【返程托运提醒】」），
# 但名字又不含交通/采购关键词，nonspot_rules 的默认判定会把它们当景点。
EXTRA_NONSPOT = ["提示", "提醒", "说明", "须知", "注意"]


def refine_kind(name, kind):
    """把上述「伪景点」从 spot 降级为事务节点。"""
    if kind == 'spot' and any(k in (name or '') for k in EXTRA_NONSPOT):
        return 'chore'
    return kind


BASE = os.path.dirname(os.path.abspath(__file__))
XLSX = os.path.join(BASE, "贵州7天6晚自驾攻略_10.1-10.7.xlsx")
OUT = os.path.join(BASE, "roadbook", "content.json")

# 景点名（带「区域·」前缀）→ (图片 slug, 地图 slug)
# 有序：先长后短，避免「黄果树景区」抢走「黄果树大瀑布」
SPOT_RULES = [
    ("龙洞堡",     (None,          "longdongbao")),
    ("黔灵山",     ("qianlingshan", "qianlingshan")),
    ("青云市集",   ("qingyunshiji", "qingyunshiji")),
    ("黄果树大瀑布", ("huangguoshu", "huangguoshu")),
    ("黄果树景区", ("huangguoshu",  "huangguoshu")),
    ("陡坡塘",     ("doupotang",    "doupotang")),
    ("天星桥",     ("tianxingqiao", "tianxingqiao")),
    ("小七孔东门", ("xiaoqikong",   "xiaoqikong")),
    ("荔波小七孔", ("xiaoqikong",   "xiaoqikong")),
    ("小七孔",     ("xiaoqikong",   "xiaoqikong")),
    ("卧龙潭",     ("wolongtan",    "xiaoqikong")),
    ("荔波县城",   ("libo_city",    "libo_city")),
    ("荔波",       ("libo_city",    "libo_city")),
    ("西江",       ("xijiang",      "xijiang")),
    ("青岩",       ("qingyan",      "qingyan")),
    ("甲秀楼",     ("jiaxiulou",    "jiaxiulou")),
    ("贵阳市区",   ("guiyang_city", "guiyang_city")),
    ("贵阳",       (None,           "guiyang_city")),
]


def match_spot(name):
    for kw, (img, mp) in SPOT_RULES:
        if kw in name:
            return img, mp
    return None, None
CITY_MAP = {1: "guiyang_city", 2: "anshun_city", 3: "libo_city",
            4: "libo_city", 5: "qdn_city", 6: "guiyang_city", 7: "guiyang_city"}
CITY_NAME = {1: "贵阳", 2: "安顺·黄果树", 3: "荔波", 4: "荔波",
             5: "黔东南·西江", 6: "贵阳", 7: "贵阳"}


# ⚠️ [已废弃] 只认「【】」的旧实现：会把 〔交通〕/〔事项〕 块当成「前言块」，
# 导致机场 / 取车点 / 特产采购这类内容被静默丢弃或错误并入上一个景点。
# 主流程已改用文件顶部 import 的共享模块 nonspot_rules.split_blocks（返回带 kind）。
# 改名 legacy_ 前缀是为了不遮蔽那份 import —— 请勿在新代码中调用。
def legacy_split_blocks(text):
    """把「【景点】\\n详情…\\n\\n【景点】\\n…」拆成 [(name, body)]"""
    if not text:
        return []
    blocks = []
    # 以【…】为分隔
    parts = re.split(r'(?=【[^】]{1,20}】)', text)
    for p in parts:
        p = p.strip()
        if not p:
            continue
        m = re.match(r'【([^】]{1,20})】(.*)', p, re.S)
        if m:
            blocks.append((m.group(1).strip(), m.group(2).strip()))
        else:
            blocks.append((None, p))
    return blocks


def lines_of(text):
    return [l.strip() for l in (text or '').split('\n') if l.strip()]


def core(name):
    """去掉「区域·」前缀与括号补充，得到景点核心名。"""
    if not name:
        return ''
    s = name.split('·')[-1]
    s = re.sub(r'[（(].*?[)）]', '', s)
    s = re.sub(r'\s+', '', s)
    return s.strip()


def match_pick(blocks, sec, body_key=''):
    """把 H/I 列的块挂到 D 列景点上：核心名互相包含，或块正文提到该景点名。"""
    c = core(sec['name'])
    short = core(sec.get('short', ''))
    hit = []
    for n, b in blocks:
        if not n:
            continue
        cn = core(n)
        if cn and c and (cn in c or c in cn):
            hit.append((n, b))
            continue
        if short and len(short) >= 2 and short in b:
            hit.append((n, b))
    return hit


def parse_day(ws, r):
    a = (ws.cell(r, 1).value or '').replace('\u3000', ' ')
    dm = re.match(r'DAY(\d+)\s*([\d/]+)?\s*(\S*)?', a)
    day = int(dm.group(1)) if dm else r - 1
    date = (dm.group(2) or '').strip() if dm else ''
    weekday = (dm.group(3) or '').strip() if dm else ''

    b_all = (ws.cell(r, 2).value or '').strip()
    b_lines = [l.strip() for l in b_all.split('\n') if l.strip()]
    stay_note = ''
    b_main = ''
    for l in b_lines:
        if l.startswith('（') or l.startswith('('):
            stay_note = l.strip('（）() ')
            continue
        if not b_main:
            b_main = l
    b_main = re.sub(r'[（(].*?[)）]', '', b_main).strip()
    # 模板里的占位说明不是真实内容，剔除
    if '出发地城市' in stay_note or '住宿城市→' in stay_note:
        stay_note = ''
    stay_note = stay_note.replace('住宿城市：', '').strip()
    parts = re.split(r'[→✈➔➜]|->', b_main)
    parts = [p.strip() for p in parts if p.strip()]
    frm = parts[0] if parts else b_main
    to = parts[-1] if len(parts) > 1 else ''

    brief = lines_of(ws.cell(r, 3).value)
    food = lines_of(ws.cell(r, 6).value)
    hotel = lines_of(ws.cell(r, 7).value)

    day_shots, day_tips, day_notice = [], [], []
    sections = []

    # ---- D 列：景点建块；〔交通〕/〔事项〕/前言块降级为当日注意事项 ----
    for name, body, kind in split_blocks(ws.cell(r, 4).value):
        kind = refine_kind(name, kind)
        if kind == 'spot':
            body, tail = split_tail(body)
            img_slug, map_slug = match_spot(name)
            sections.append(dict(
                name=name, short=core(name),
                slug=img_slug, map=map_slug,
                detail=body, shots=[], tips=[], notice=[],
            ))
            for t in tail:
                day_tips.append(tail_block(t))
        else:
            day_notice.append(dict(name=block_label(name, kind),
                                   lines=clean_items(lines_of(body)), kind=kind))

    # ---- E 列：机位。景点块挂景点；前言块作日级机位；非景点块整块丢弃 ----
    for name, body, kind in split_blocks(ws.cell(r, 5).value):
        ls = clean_items(lines_of(body))
        if not ls or kind in ('transit', 'chore'):
            continue                       # 机场 / 取车 / 采购点没有「机位」概念
        placed = False
        if name:
            cn = core(name)
            for s in sections:
                if cn and (cn in s['short'] or s['short'] in cn):
                    s['shots'].append(dict(name=name, lines=ls))
                    placed = True
                    break
        if not placed:
            day_shots.append(dict(name=name or '当日机位', lines=ls))

    # ---- H 列：景点块挂景点；非景点块 → 当日备注（以 🚉/🛍 节点名起小标题）----
    for name, body, kind in split_blocks(ws.cell(r, 8).value):
        kind = refine_kind(name, kind)
        body, tail = split_tail(body)
        ls = lines_of(body)
        if kind == 'spot':
            cn = core(name)
            hit = next((s for s in sections
                        if cn and (cn in s['short'] or s['short'] in cn)), None)
            if hit:
                hit['tips'].extend(ls)
            else:
                day_tips.append(dict(name=name, lines=ls))
        else:
            day_tips.append(dict(name=block_label(name, kind), lines=ls))
        for t in tail:
            day_tips.append(tail_block(t))

    # ---- I 列：景点块挂景点；其余全部并入当日注意事项 ----
    for name, body, kind in split_blocks(ws.cell(r, 9).value):
        kind = refine_kind(name, kind)
        body, tail = split_tail(body)
        ls = lines_of(body)
        cn = core(name) if name else ''
        hit = next((s for s in sections
                    if cn and (cn in s['short'] or s['short'] in cn)), None)
        if kind == 'spot' and hit:
            hit['notice'].extend(ls)
        else:
            day_notice.append(dict(name=block_label(name, kind), lines=ls, kind=kind))
        for t in tail:
            day_notice.append(tail_block(t))

    # 机位：每个大景点下如果一条都没有，就把日级机位并到唯一 section 上
    if len(sections) == 1 and day_shots:
        sections[0]['shots'].extend(day_shots)
        day_shots = []

    return dict(day=day, date=date, weekday=weekday,
                from_city=frm, to_city=to, stay_note=stay_note,
                transport_raw=b_main,
                brief=brief, food=food, hotel=hotel,
                sections=sections,
                day_shots=day_shots, day_tips=day_tips, day_notice=day_notice,
                city=CITY_NAME.get(day, ''), city_map=CITY_MAP.get(day))


def parse_budget(wb):
    ws = wb['人均预算']
    title = note = warning = ''
    rows, totals = [], []
    per_weekday = per_holiday = tot_weekday = tot_holiday = delta = ''
    for r in range(1, ws.max_row + 1):
        v = [ws.cell(r, c).value for c in range(1, 7)]
        s = [str(x).strip() if x is not None else '' for x in v]
        joined = ' '.join(s)
        if '口径说明' in s[0]:
            note = s[0]
        if s[0].startswith('⚠️') or '不确定性' in s[0]:
            warning = s[0]
        if s[0] == '类别':
            headers = s
            continue
        is_total = ('合计' in s[2]) or (s[2].count('¥') and '/ 人' in s[2])
        if is_total:
            if '平日合计' in s[2]:
                continue
            totals.append(s)
            per_weekday = s[2].replace(' / 人', '').strip()
            per_holiday = s[3].replace(' / 人', '').strip()
            delta = s[4]
            m = re.search(r'平日\s*¥([\d,]+)', s[5])
            if m:
                tot_weekday = '¥' + m.group(1)
            m = re.search(r'国庆\s*¥([\d,]+)', s[5])
            if m:
                tot_holiday = '¥' + m.group(1)
            continue
        if s[2].startswith('¥') and s[3].startswith('¥'):
            rows.append(s)
    return dict(title=title, note=note, rows=rows, totals=totals, warning=warning,
                headers=headers if 'headers' in dir() else ['类别', '计价口径', '平日参考价（人均）', '国庆实际价（人均）', '涨幅', '数据来源与说明'],
                per_weekday=per_weekday, per_holiday=per_holiday,
                total_weekday=tot_weekday, total_holiday=tot_holiday, delta=delta)


def main():
    wb = load_workbook(XLSX)
    ws = wb['全部行程']
    days = [parse_day(ws, r) for r in range(2, ws.max_row + 1)]
    days = [d for d in days if d['sections']]

    budget = parse_budget(wb)

    ws_map = wb['景点地图']
    maps = []
    for r in range(2, ws_map.max_row + 1):
        row = [ws_map.cell(r, c).value for c in range(1, 6)]
        if row[1]:
            maps.append(dict(type=str(row[0] or ''), name=str(row[1]),
                             kw=str(row[2] or ''), link=str(row[3] or ''),
                             note=str(row[4] or '')))

    ws_pack = wb['物品清单']
    packing = []
    for r in range(1, ws_pack.max_row + 1):
        row = [ws_pack.cell(r, c).value for c in range(1, 8)]
        packing.append([str(x or '') for x in row])

    content = dict(
        meta=dict(
            title="贵州 7 天 6 晚 自驾旅游攻略",
            destination="贵州", days=7, nights=6, mode="自驾",
            dateRange="2026.10.1 - 10.7", people=2, gender="2 男",
            departure="深圳", transport="飞机往返",
            pace="舒适型 · 每日 9:00 出发", level="经济型",
            route="贵阳环线·黔中铁三角",
            stops=["贵阳", "黄果树", "荔波", "西江", "贵阳"],
            miles="约 900 km",
        ),
        budget=budget,
        days=days,
        maps=maps,
        packing=packing,
    )
    with open(OUT, 'w', encoding='utf-8') as f:
        json.dump(content, f, ensure_ascii=False, indent=1)
    print('written', OUT, os.path.getsize(OUT))
    print('days:', [(d['day'], len(d['sections']), [s['name'] for s in d['sections']]) for d in days])
    print('budget rows:', len(budget['rows']), '| totals:', len(budget['totals']))
    print('maps:', len(maps), '| packing rows:', len(packing))


if __name__ == '__main__':
    main()
